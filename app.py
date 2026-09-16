
import io, re, hashlib, json, sqlite3
from datetime import datetime
from pathlib import Path
import streamlit as st
from pypdf import PdfReader, PdfWriter
from docx import Document

DB="score_engine.db"
st.set_page_config(page_title="PLEADS SCORE ENGINE V2",page_icon="📊",layout="wide")

def db():
    con=sqlite3.connect(DB); con.row_factory=sqlite3.Row
    con.executescript("""
    CREATE TABLE IF NOT EXISTS projects(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT NOT NULL,created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS uploads(id INTEGER PRIMARY KEY AUTOINCREMENT,project_id INTEGER,filename TEXT,sha256 TEXT UNIQUE,judge TEXT,uploaded_at TEXT,raw BLOB,analysis_json TEXT);
    CREATE TABLE IF NOT EXISTS verifications(project_id INTEGER,team TEXT,status TEXT DEFAULT 'Pending',note TEXT DEFAULT '',updated_at TEXT,PRIMARY KEY(project_id,team));
    """)
    return con
con=db()

def clean(s): return re.sub(r"\s+"," ",(s or "").replace("\u00a0"," ")).strip()
def norm(s): return re.sub(r"\s+"," ",re.sub(r"[^a-z0-9\s.]","",clean(s).lower())).strip()
def canonical(s,aliases):
    s=clean(s)
    for k,v in aliases.items():
        if norm(k)==norm(s): return v
    return s
def nums(s): return [int(x) for x in re.findall(r"(?<!\d)(\d{1,4})(?!\d)",s or "")]
def filehash(b): return hashlib.sha256(b).hexdigest()

def pages_text(data):
    r=PdfReader(io.BytesIO(data))
    return ["\n".join(clean(x) for x in (p.extract_text() or "").splitlines() if clean(x)) for p in r.pages]

def detect_judge(text,filename):
    # The completed PDF often has a blank Juri field; filename is the reliable
    # source in this workflow: [Yasin] ..., [Marcell] ..., etc.
    for pat in [r"(?im)\bjuri\s*:\s*([^\n]+)", r"(?im)\bjuri\s*[-:]\s*([^\n]+)"]:
        m=re.search(pat,text)
        if m:
            val=clean(m.group(1))
            val=re.split(r"\bJudul\s+Berkas\s*:",val,flags=re.I)[0].strip()
            if val and val.lower() not in {"nama","-"}: return val
    m=re.search(r"^\[([^\]]+)\]",Path(filename).stem)
    if m: return clean(m.group(1))
    m=re.search(r"^\(([^)]+)\)",Path(filename).stem)
    return clean(m.group(1)) if m else ""

def team_for_page(text,previous,aliases):
    lines=[clean(x) for x in text.splitlines() if clean(x)]
    for i,line in enumerate(lines):
        m=re.search(r"\bnama\s*tim\s*:\s*(.*)$",line,re.I)
        if m:
            val=clean(m.group(1))
            # If the value is on the same extracted line, stop before boilerplate.
            val=re.split(r"\bPenentuan\s+pemenang\b",val,flags=re.I)[0].strip()
            if val.lower() not in {"","-","nama tim"}: return canonical(val,aliases)
            # Some PDF extractors put the value on the next line.
            if i+1<len(lines):
                val=clean(lines[i+1])
                if val and val.lower() not in {"-","nama tim"}: return canonical(val,aliases)
    for line in lines:
        m=re.search(r"^\s*[Pp]\s*[-–]\s*(.+)$",line)
        if m: return canonical(m.group(1),aliases)
        m=re.search(r"^\s*\d+\s*[-–]\s+(.+)$",line)
        if m and len(clean(m.group(1)))<100: return canonical(m.group(1),aliases)
    return previous

def schema(text):
    lines=[clean(x) for x in text.splitlines() if clean(x)]
    out=[]; seen=set()
    scale_re=re.compile(r"^0\s*[-–]\s*(\d{2,4})(?:\s+(\d{1,4}))?$")
    for i,line in enumerate(lines):
        m=scale_re.match(line)
        if not m: continue
        mx=int(m.group(1)); name=""
        for j in range(i-1,max(-1,i-4),-1):
            cand=clean(lines[j])
            if re.fullmatch(r"\d+",cand) or cand.lower() in {"no. aspek penilaian skala skor","no aspek penilaian bobot skor"}: continue
            # Aspect rows may contain their subcriteria on the same line.
            cand=re.split(r"\s+(?=[a-z]{1,3}\.\s)",cand,maxsplit=1,flags=re.I)[0].strip()
            if len(cand)>=3 and len(cand)<=100 and not re.search(r"(penentuan pemenang|lembar penilaian|komentar|total 1000)",cand,re.I):
                name=re.sub(r"^\d+\s+", "", cand).strip(); break
        if name and norm(name) not in seen and norm(name) not in {"aspek penilaian","skala","skor","bobot"}:
            out.append((name,mx)); seen.add(norm(name))
    return out

def scores(text,aspects):
    lines=[clean(x) for x in text.splitlines() if clean(x)]
    vals={}; scale_re=re.compile(r"^0\s*[-–]\s*(\d{2,4})(?:\s+(\d{1,4}))?$")
    for a,mx in aspects:
        for i,line in enumerate(lines):
            if norm(a) in norm(line):
                for j in range(i+1,min(i+12,len(lines))):
                    sm=scale_re.match(lines[j])
                    if sm and int(sm.group(1))==mx:
                        if sm.group(2) is not None: vals[a]=int(sm.group(2))
                        else:
                            for k in range(j+1,min(j+3,len(lines))):
                                ns=[n for n in nums(lines[k]) if n<=mx]
                                if ns: vals[a]=ns[-1]; break
                        break
                if a in vals: break
    total=None
    for i,line in enumerate(lines):
        m=re.match(r"^total\s+1000(?:\s+(\d{1,4}))?",line,re.I)
        if m:
            if m.group(1) is not None: total=int(m.group(1))
            else:
                ns=[n for n in nums(line) if n!=1000]
                if ns: total=ns[-1]
            break
    return vals,total

def comment(text):
    m=re.search(r"(?is)\bkomentar\s*:\s*(.*)$",text)
    return clean(m.group(1)) if m else ""

def analyze(data,filename,aliases):
    ts=pages_text(data); judge=detect_judge("\n".join(ts),filename); current=""; rec=[]
    for page,t in enumerate(ts,1):
        team=team_for_page(t,current,aliases)
        if team: current=team
        # For the current LOC HUKOL workflow, only completed BERKAS scoring
        # sheets count as records. Presentation sheets are a separate stage
        # and may be blank, so they must not create false records.
        is_berkas = bool(re.search(r"(?i)LEMBAR\s+PENILAIAN\s+BERKAS", t))
        if not is_berkas: continue
        asp=schema(t); sc,total=scores(t,asp)
        if team and asp:
            rec.append({"page":page,"team":team,"judge":judge,"type":"berkas","aspects":asp,"scores":sc,"total":total,"comment":comment(t)})
    return {"filename":filename,"judge":judge,"pages":len(ts),"records":rec}

def split_pdf(data,pages):
    r=PdfReader(io.BytesIO(data)); w=PdfWriter()
    for p in pages: w.add_page(r.pages[p-1])
    out=io.BytesIO(); w.write(out); return out.getvalue()

def merge_pdfs(chunks):
    w=PdfWriter()
    for _,b in chunks:
        r=PdfReader(io.BytesIO(b))
        for page in r.pages: w.add_page(page)
    out=io.BytesIO(); w.write(out); return out.getvalue()

def docx_recap(team,recs):
    d=Document(); d.add_heading(f"Rekap Penilaian — {team}",0)
    for r in recs:
        d.add_heading(r["judge"] or "Juri",1)
        d.add_paragraph(f"Sumber: {r['filename']} | Halaman: {r['page']}")
        t=d.add_table(rows=1,cols=3); t.style="Table Grid"
        for c,v in zip(t.rows[0].cells,["Aspek","Maks.","Skor"]): c.text=v
        for a,mx in r["aspects"]:
            c=t.add_row().cells; c[0].text=a; c[1].text=str(mx); c[2].text=str(r["scores"].get(a,""))
        calc=sum(r["scores"].values()) if r["scores"] else ""
        d.add_paragraph(f"Total tertulis: {r['total'] if r['total'] is not None else ''} | Jumlah aspek: {calc}")
        if r["comment"]: d.add_paragraph("Komentar: "+r["comment"])
    out=io.BytesIO(); d.save(out); return out.getvalue()

def validation(analyses,expected):
    issues=[]
    for a in analyses:
        local={r["team"] for r in a["records"] if r["team"]}
        if not a["judge"]: issues.append(("error",a["filename"],"Juri belum terdeteksi"))
        if not a["records"]: issues.append(("error",a["filename"],"Tidak ada lembar penilaian terdeteksi"))
        for r in a["records"]:
            missing=[x for x,_ in r["aspects"] if x not in r["scores"]]
            calc=sum(r["scores"].values()) if r["scores"] else None
            if missing: issues.append(("warning",a["filename"],f"{r['team']}: skor tidak terbaca pada {', '.join(missing)}"))
            if r["total"] is not None and calc is not None and r["total"]!=calc:
                issues.append(("error",a["filename"],f"{r['team']}: total {r['total']} ≠ jumlah aspek {calc}"))
        if expected:
            miss=set(expected)-local; extra=local-set(expected)
            if miss: issues.append(("warning",a["filename"],"Tim tidak ditemukan: "+", ".join(sorted(miss))))
            if extra: issues.append(("warning",a["filename"],"Tim ekstra/beda nama: "+", ".join(sorted(extra))))
    return issues

st.markdown("# 📊 PLEADS SCORE ENGINE V2")
st.caption("Upload → auto-detect → check → pisah → gabung → verifikasi → publication gate")

with st.sidebar:
    st.header("Project")
    ps=con.execute("SELECT * FROM projects ORDER BY id DESC").fetchall()
    opts=["+ Project Baru"]+[f"{p['id']} — {p['name']}" for p in ps]
    choice=st.selectbox("Pilih project",opts)
    if choice=="+ Project Baru":
        nm=st.text_input("Nama project",placeholder="Seleksi Internal LOC HUKOL 2026")
        if st.button("Buat project",type="primary") and nm.strip():
            con.execute("INSERT INTO projects(name,created_at) VALUES(?,?)",(nm.strip(),datetime.now().isoformat())); con.commit(); st.rerun()
        pid=None
    else: pid=int(choice.split(" — ",1)[0])
    st.divider(); st.subheader("Alias nama tim")
    ar=st.text_area("salah = benar",placeholder="Karl Max = Karl Marx")
    aliases={}
    for line in ar.splitlines():
        if "=" in line:
            k,v=line.split("=",1); aliases[k.strip()]=v.strip()

if not pid:
    st.info("Buat project dulu.")
    st.stop()

ups=con.execute("SELECT * FROM uploads WHERE project_id=? ORDER BY id",(pid,)).fetchall()
analyses=[json.loads(x["analysis_json"]) for x in ups]
raw={x["id"]:x["raw"] for x in ups}

tabs=st.tabs(["1 · UPLOAD","2 · CHECK","3 · PISAH","4 · GABUNG","5 · VERIFIKASI","6 · PUBLIKASI"])

with tabs[0]:
    fs=st.file_uploader("Upload satu atau banyak PDF",type="pdf",accept_multiple_files=True)
    if fs and st.button("🔍 Analisis & simpan",type="primary"):
        for f in fs:
            b=f.getvalue(); h=filehash(b)
            if con.execute("SELECT 1 FROM uploads WHERE sha256=?",(h,)).fetchone():
                st.warning("Skip duplikat: "+f.name); continue
            a=analyze(b,f.name,aliases)
            con.execute("INSERT INTO uploads(project_id,filename,sha256,judge,uploaded_at,raw,analysis_json) VALUES(?,?,?,?,?,?,?)",(pid,f.name,h,a["judge"],datetime.now().isoformat(),b,json.dumps(a,ensure_ascii=False)))
        con.commit(); st.success("Selesai."); st.rerun()
    st.dataframe([{"File":x["filename"],"Juri":x["judge"] or "⚠️ belum terdeteksi"} for x in ups],use_container_width=True,hide_index=True)

with tabs[1]:
    if not analyses: st.info("Upload dulu.")
    else:
        auto=sorted({r["team"] for a in analyses for r in a["records"] if r["team"]})
        exp=[clean(x) for x in st.text_area("Daftar tim seharusnya (satu per baris)",value="\n".join(auto)).splitlines() if clean(x)]
        iss=validation(analyses,exp)
        a,b,c=st.columns(3); a.metric("File juri",len(analyses)); b.metric("Tim",len(auto)); c.metric("Issue",len(iss))
        for typ,file,msg in iss: st.write(("🔴" if typ=="error" else "🟠")+f" **{file}** — {msg}")
        rows=[]
        for x in analyses:
            for r in x["records"]:
                calc=sum(r["scores"].values()) if r["scores"] else None
                rows.append({"Juri":r["judge"],"Tim":r["team"],"Hal.":r["page"],"Aspek":len(r["aspects"]),"Skor":len(r["scores"]),"Total":r["total"],"Hitung":calc,"Status":"OK" if len(r["scores"])==len(r["aspects"]) and (r["total"] is None or r["total"]==calc) else "CHECK"})
        st.dataframe(rows,use_container_width=True,hide_index=True)

with tabs[2]:
    teams=sorted({r["team"] for a in analyses for r in a["records"] if r["team"]})
    if teams:
        t=st.selectbox("Tim",teams)
        for u in ups:
            a=json.loads(u["analysis_json"]); rs=[r for r in a["records"] if r["team"]==t]
            if rs:
                st.download_button(f"⬇️ {a['judge'] or 'Juri'} — {t}",split_pdf(raw[u["id"]],[r["page"] for r in rs]),file_name=f"{t} — {a['judge'] or 'Juri'}.pdf",mime="application/pdf",key=f"s{u['id']}{norm(t)}")
        rs=[r for a in analyses for r in a["records"] if r["team"]==t]
        st.download_button("⬇️ Rekap DOCX per tim",docx_recap(t,rs),file_name=f"{t} — Rekap Penilaian.docx",mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    else: st.info("Belum ada tim.")

with tabs[3]:
    teams=sorted({r["team"] for a in analyses for r in a["records"] if r["team"]})
    for t in teams:
        chunks=[]
        for u in ups:
            a=json.loads(u["analysis_json"]); rs=[r for r in a["records"] if r["team"]==t]
            if rs: chunks.append((a["judge"],split_pdf(raw[u["id"]],[r["page"] for r in rs])))
        if chunks:
            st.download_button(f"⬇️ {t} — MERGED ({len(chunks)} juri)",merge_pdfs(chunks),file_name=f"{t} — MERGED.pdf",mime="application/pdf",key="m"+norm(t))

with tabs[4]:
    teams=sorted({r["team"] for a in analyses for r in a["records"] if r["team"]})
    for t in teams:
        row=con.execute("SELECT * FROM verifications WHERE project_id=? AND team=?",(pid,t)).fetchone()
        old=row["status"] if row else "Pending"; note=row["note"] if row else ""
        c1,c2,c3=st.columns([2,2,4]); c1.write("**"+t+"**")
        status=c2.selectbox("Status",["Pending","Terkirim","Confirmed","Revisi"],index=["Pending","Terkirim","Confirmed","Revisi"].index(old),key="vs"+norm(t))
        nt=c3.text_input("Catatan",value=note,key="vn"+norm(t))
        if status!=old or nt!=note:
            con.execute("INSERT INTO verifications(project_id,team,status,note,updated_at) VALUES(?,?,?,?,?) ON CONFLICT(project_id,team) DO UPDATE SET status=excluded.status,note=excluded.note,updated_at=excluded.updated_at",(pid,t,status,nt,datetime.now().isoformat())); con.commit()

with tabs[5]:
    teams=sorted({r["team"] for a in analyses for r in a["records"] if r["team"]})
    iss=validation(analyses,teams)
    ver={x["team"]:x["status"] for x in con.execute("SELECT team,status FROM verifications WHERE project_id=?",(pid,)).fetchall()}
    pending=[t for t in teams if ver.get(t)!="Confirmed"]
    if iss: st.error(f"{len(iss)} issue masih harus dibereskan.")
    if pending: st.warning("Belum Confirmed: "+", ".join(pending))
    if teams and not iss and not pending: st.success("✅ PUBLICATION GATE LULUS")
    rows=[]
    for t in teams:
        rs=[r for a in analyses for r in a["records"] if r["team"]==t]
        vals=[(r["judge"],r["total"] if r["total"] is not None else (sum(r["scores"].values()) if r["scores"] else None)) for r in rs]
        vals=[x for x in vals if x[1] is not None]
        rows.append({"Tim":t,"Jumlah Juri":len(vals),"Rata-rata":sum(v for _,v in vals)/len(vals) if vals else None,"Rincian":"; ".join(f"{j}: {v}" for j,v in vals)})
    st.dataframe(rows,use_container_width=True,hide_index=True)
    csv="Tim,Jumlah Juri,Rata-rata,Rincian\n"+"\n".join(f'"{r["Tim"]}",{r["Jumlah Juri"]},{r["Rata-rata"] if r["Rata-rata"] is not None else ""},"{r["Rincian"].replace(chr(34),chr(34)*2)}"' for r in rows)
    st.download_button("⬇️ Export Central Recap CSV",csv.encode("utf-8-sig"),file_name="central_recap.csv",mime="text/csv")
